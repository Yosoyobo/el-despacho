"""La Carga Contable (S-Carga-Contable): la contabilidad de fuera, de un jalón.

Se prueba con la plantilla DE VERDAD: se genera con `plantilla.generar()`, se
llena con openpyxl como lo haría una persona y se sube. Cubre:

- La plantilla trae sus hojas, listas y cuentas precargadas.
- La lectura tolera fechas y montos como los pega una persona.
- La vista previa no guarda nada; aplicar guarda exactamente lo previsto.
- El cuadre: apertura y «saldos hoy» registran sólo la diferencia.
- Lo ya capturado se reconoce (±3 días, mismo monto) y no se duplica;
  volver a subir el mismo archivo tampoco.
- Lo anterior al arranque (facturas por cobrar, gastos por pagar) no toca el año.
- Un cobro se liga solo a su factura cuando es inequívoco.
- Con un error no se aplica nada.
- No salen correos a clientes.
- Deshacer deja los saldos como estaban y libera los folios.
- Permisos: `contaduria.cargar`.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest

pytestmark = [pytest.mark.django_db, pytest.mark.taller]

ARRANQUE = date(2026, 1, 1)
HOY = date(2026, 9, 29)


# ── Utilidades ───────────────────────────────────────────────────────────

@pytest.fixture(autouse=True)
def _chalan_sin_red(monkeypatch):
    """Nunca se llama a un proveedor real. Por default El Chalán «no contesta»
    (la carga sigue con las reglas de siempre); las pruebas de IA lo sustituyen."""
    from apps.contaduria.carga import ia

    def _no_contesta(prompt, usuario):
        raise RuntimeError("sin red en pruebas")
    monkeypatch.setattr(ia, "_llamar", _no_contesta)


def _chalan_dice(monkeypatch, respuesta_por_descripcion: dict):
    """El Chalán contesta según la descripción de cada movimiento del prompt."""
    import json
    import re

    from apps.contaduria.carga import ia
    preguntas = []

    def _llamar(prompt, usuario):
        preguntas.append(prompt)
        movs = json.loads(re.search(r"MOVIMIENTOS:\n(\[.*\])", prompt, re.DOTALL).group(1))
        salida = []
        for m in movs:
            for clave, resp in respuesta_por_descripcion.items():
                if clave in (m.get("descripcion") or ""):
                    salida.append({"id": m["id"], **resp})
        return json.dumps({"movimientos": salida})
    monkeypatch.setattr(ia, "_llamar", _llamar)
    return preguntas


@pytest.fixture
def centro(db):
    from apps.tesoreria.models import CentroDeCosto
    return CentroDeCosto.objects.create(nombre="Operación general", slug="operacion", naturaleza="operativo")


@pytest.fixture
def actor(usuario_factory):
    return usuario_factory(rol="super_admin", email="oscar@ejemplo.com")


def _llenar(*, arranque=None, fecha_arranque=ARRANQUE, hoy=None, fecha_hoy=HOY,
            ingresos=(), gastos=(), facturas=(), polizas=()) -> bytes:
    """Genera la plantilla real y la llena. Cada renglón es un dict por clave
    de columna del esquema; los saldos, {etiqueta o código de cuenta: saldo}."""
    from apps.contaduria.carga import esquema as E
    from apps.contaduria.carga import plantilla
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(plantilla.generar()))

    def saldos(nombre, fecha, valores):
        ws = wb[nombre]
        ws[E.CELDA_FECHA] = fecha
        valores = dict(valores or {})
        fila = E.FILA_ENCABEZADO_SALDOS + 1
        while ws.cell(row=fila, column=1).value:
            etiqueta = ws.cell(row=fila, column=1).value
            codigo = etiqueta.split(" ")[0]
            for k in (etiqueta, codigo):
                if k in valores:
                    ws.cell(row=fila, column=2, value=valores.pop(k))
            fila += 1
        for cuenta, saldo in valores.items():   # cuentas no precargadas
            ws.cell(row=fila, column=1, value=cuenta)
            ws.cell(row=fila, column=2, value=saldo)
            fila += 1

    def renglones(nombre, filas):
        columnas = E.HOJAS_RENGLONES[nombre]
        ws = wb[nombre]
        for i, datos in enumerate(filas):
            for j, col in enumerate(columnas, start=1):
                if col.clave in datos:
                    ws.cell(row=E.FILA_ENCABEZADO + 1 + i, column=j, value=datos[col.clave])

    saldos(E.HOJA_ARRANQUE, fecha_arranque, arranque)
    saldos(E.HOJA_HOY, fecha_hoy, hoy)
    renglones(E.HOJA_INGRESOS, ingresos)
    renglones(E.HOJA_GASTOS, gastos)
    renglones(E.HOJA_FACTURAS, facturas)
    renglones(E.HOJA_POLIZAS, polizas)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _saldo(slot_o_codigo: str) -> Decimal:
    from apps.contaduria.models import CuentaContable
    from apps.contaduria.services import saldo_cuenta

    c = (CuentaContable.objects.filter(slot=slot_o_codigo).first()
         or CuentaContable.objects.get(codigo=slot_o_codigo))
    return saldo_cuenta(c)


def _nueva(contenido, actor, estados=(), crear=True):
    from apps.contaduria.carga import motor
    return motor.nueva(plantilla=contenido, nombre_plantilla="prueba.xlsx", estados=estados,
                       crear_desde_estados=crear, actor=actor)


def _banco():
    from apps.contaduria.models import CuentaContable
    return CuentaContable.objects.get(slot="banco")


def _aplicar(contenido, actor, estados=()):
    from apps.contaduria.carga import motor
    carga = _nueva(contenido, actor, estados)
    assert carga.resumen["errores"] == 0, [p for p in carga.resumen["pasos"] if p["estado"] == "error"]
    motor.aplicar(carga, actor)
    carga.refresh_from_db()
    return carga


def _estado(carga, hoja):
    return [p["estado"] for p in carga.resumen["pasos"] if p["hoja"] == hoja]


# ── Plantilla y lectura ──────────────────────────────────────────────────

def test_plantilla_trae_hojas_listas_y_cuentas(centro, cliente_factory):
    from apps.contaduria.carga import esquema as E
    from apps.contaduria.carga import plantilla
    from openpyxl import load_workbook

    cliente_factory(razon_social="Heladería La Fuente")
    wb = load_workbook(io.BytesIO(plantilla.generar()))
    for nombre in (E.HOJA_LEEME, E.HOJA_ARRANQUE, E.HOJA_INGRESOS, E.HOJA_GASTOS,
                   E.HOJA_FACTURAS, E.HOJA_POLIZAS, E.HOJA_HOY, E.HOJA_LISTAS):
        assert nombre in wb.sheetnames
    assert wb.sheetnames[0] == E.HOJA_LEEME
    assert wb[E.HOJA_LISTAS].sheet_state == "hidden"
    precargadas = [c.value for c in wb[E.HOJA_ARRANQUE]["A"][E.FILA_ENCABEZADO_SALDOS:] if c.value]
    assert any(v.startswith("1.1.02") for v in precargadas)          # Bancos sí
    assert not any(v.startswith("1.2.01") for v in precargadas)      # Clientes sale del detalle
    listas = wb[E.HOJA_LISTAS]
    todas = [c.value for col in listas.iter_cols() for c in col if c.value]
    assert "Heladería La Fuente" in todas and "Operación general" in todas
    assert wb[E.HOJA_INGRESOS].data_validations.dataValidation  # hay desplegables


def test_lectura_tolera_lo_que_pega_una_persona():
    from apps.contaduria.carga import lectura

    assert lectura.leer_fecha("3/2/26") == date(2026, 2, 3)
    assert lectura.leer_fecha("2026-02-03") == date(2026, 2, 3)
    assert lectura.leer_fecha(46056) == date(2026, 2, 3)
    assert lectura.leer_fecha("") is None
    assert lectura.leer_monto("$1,200.50") == Decimal("1200.50")
    assert lectura.leer_monto("(300)") == Decimal("-300.00")
    assert lectura.leer_monto("") is None
    assert lectura.leer_monto(0) == Decimal("0.00")   # 0 es cero, no «vacío»
    assert lectura.leer_si_no("sí") is True and lectura.leer_si_no("") is False
    with pytest.raises(ValueError):
        lectura.leer_fecha("mañana")


def test_archivo_que_no_es_la_plantilla(actor):
    from apps.contaduria.carga import lectura, motor

    with pytest.raises(lectura.PlantillaInvalida):
        motor.nueva(plantilla=b"no soy un excel", nombre_plantilla="x.xlsx", actor=actor)
    with pytest.raises(lectura.PlantillaInvalida):
        motor.nueva(actor=actor)  # nada que cargar


# ── Vista previa vs aplicar ──────────────────────────────────────────────

def test_vista_previa_no_guarda_nada(centro, actor):
    from apps.contaduria.models import Asiento
    from apps.tesoreria.models import Egreso, Ingreso

    contenido = _llenar(arranque={"1.1.02": 10000},
                        ingresos=[{"fecha": date(2026, 2, 1), "descripcion": "Pago", "monto": 5000}],
                        gastos=[{"fecha": date(2026, 2, 3), "descripcion": "Tinta", "monto": 1000}])
    antes = (Ingreso.objects.count(), Egreso.objects.count(), Asiento.objects.count())
    carga = _nueva(contenido, actor)
    assert (Ingreso.objects.count(), Egreso.objects.count(), Asiento.objects.count()) == antes
    assert carga.estado == "borrador"
    r = carga.resumen
    assert r["errores"] == 0
    assert _estado(carga, "2 Ingresos") == ["crear"] and _estado(carga, "3 Gastos") == ["crear"]
    banco = next(s for s in r["saldos"] if s["codigo"] == "1.1.02")
    assert Decimal(banco["despues"]) == Decimal("14000.00")  # 10,000 + 5,000 − 1,000
    assert Decimal(r["apertura"]["lineas"][0]["diferencia"]) == Decimal("10000.00")


def test_aplicar_cuadra_contra_el_banco(centro, actor):
    from apps.contaduria.models import Asiento

    contenido = _llenar(
        arranque={"1.1.02": 10000},
        ingresos=[{"fecha": date(2026, 2, 1), "descripcion": "Pago de mostrador", "monto": 5000}],
        gastos=[{"fecha": date(2026, 2, 3), "descripcion": "Tinta", "monto": 1000}],
        hoy={"1.1.02": 13500},  # faltan $500 que nadie anotó
    )
    carga = _aplicar(contenido, actor)
    assert carga.estado == "aplicada"
    assert _saldo("banco") == Decimal("13500.00")
    assert _saldo("3.2.01") == Decimal("10000.00")  # lo acumulado antes del arranque
    assert Asiento.vigentes.filter(origen="apertura").count() == 1
    cuadre = Asiento.vigentes.get(origen="ajuste", referencia_externa__startswith=f"contaduria.carga:{carga.pk}:cuadre")
    assert cuadre.fecha == HOY
    assert Decimal(carga.resumen["cuadre"]["lineas"][0]["diferencia"]) == Decimal("-500.00")


def test_la_apertura_solo_registra_la_diferencia(centro, actor):
    """Si El Despacho ya tenía movimientos antes del arranque, declarar el
    saldo no los cuenta dos veces."""
    from apps.contaduria.services import crear_asiento

    crear_asiento(descripcion="Previo", fecha=date(2025, 12, 15), idempotente=False,
                  partidas=[{"cuenta": "1.1.02", "cargo": 3000}, {"cuenta": "3.1.01", "abono": 3000}])
    carga = _aplicar(_llenar(arranque={"1.1.02": 10000}), actor)
    assert _saldo("banco") == Decimal("10000.00")
    linea = carga.resumen["apertura"]["lineas"][0]
    assert (linea["sistema"], linea["diferencia"]) == ("3000.00", "7000.00")


# ── Duplicados ───────────────────────────────────────────────────────────

def test_lo_ya_capturado_se_reconoce(centro, actor):
    from apps.tesoreria.models import Ingreso

    ya = Ingreso.objects.create(monto=Decimal("5000"), fecha=date(2026, 2, 2), descripcion="Capturado a mano")
    fila = {"fecha": date(2026, 2, 1), "descripcion": "Pago", "monto": 5000}
    carga = _nueva(_llenar(ingresos=[fila]), actor)
    paso = next(p for p in carga.resumen["pasos"] if p["hoja"] == "2 Ingresos")
    assert paso["estado"] == "saltar" and ya.codigo in paso["mensajes"][0]
    forzada = _nueva(_llenar(ingresos=[{**fila, "forzar": "Sí"}]), actor)
    assert _estado(forzada, "2 Ingresos") == ["crear"]


def test_subir_el_mismo_archivo_dos_veces_no_duplica(centro, actor):
    from apps.tesoreria.models import Egreso, Ingreso

    contenido = _llenar(
        arranque={"1.1.02": 10000},
        ingresos=[{"fecha": date(2026, 2, 1), "descripcion": "Pago", "monto": 5000}],
        gastos=[{"fecha": date(2026, 2, 3), "descripcion": "Café", "monto": 80},
                {"fecha": date(2026, 2, 3), "descripcion": "Café", "monto": 80}],  # dos cafés = dos gastos
    )
    _aplicar(contenido, actor)
    assert Egreso.vigentes.count() == 2
    otra = _aplicar(contenido, actor)
    assert (Ingreso.vigentes.count(), Egreso.vigentes.count()) == (1, 2)
    assert set(_estado(otra, "2 Ingresos") + _estado(otra, "3 Gastos")) == {"saltar"}
    assert otra.resumen["apertura"]["lineas"][0]["diferencia"] == "0.00"
    assert _saldo("banco") == Decimal("14840.00")


# ── Antes del arranque ───────────────────────────────────────────────────

def test_factura_anterior_al_arranque_no_toca_el_anio(centro, actor):
    from apps.contaduria.reportes import estado_resultados
    from apps.facturacion.models import Factura

    carga = _aplicar(_llenar(facturas=[{
        "folio": "F90", "cliente": "Cafetería Nueva", "rfc": "CNU010101AAA", "concepto": "Menús",
        "fecha": date(2025, 12, 10), "total": 10000, "regimen": "Exento", "cobrado_antes": 4000,
    }]), actor)
    fac = Factura.objects.get(folio_numero=90)
    assert fac.estado == "cobrada_parcial" and fac.saldo_pendiente == Decimal("6000.00")
    assert fac.cliente.razon_social == "Cafetería Nueva" and fac.cliente.rfc == "CNU010101AAA"
    assert _saldo("cxc") == Decimal("6000.00")
    assert _saldo("banco") == Decimal("0.00")                 # el cobro de antes no inventa dinero hoy
    assert _saldo("3.2.01") == Decimal("6000.00")             # lo que se debía al arrancar
    er = estado_resultados(desde=date(2026, 1, 1), hasta=HOY)
    assert er["ingresos"]["total"] == Decimal("0.00")
    assert carga.resumen["clientes_nuevos"] == ["Cafetería Nueva"]


def test_gasto_pendiente_anterior_queda_por_pagar(centro, actor):
    from apps.contaduria.reportes import estado_resultados

    carga = _aplicar(_llenar(gastos=[
        {"fecha": date(2025, 12, 20), "descripcion": "Maquila de diciembre", "monto": 2000, "estado": "Pendiente de pago"},
    ]), actor)
    assert _saldo("cxp") == Decimal("2000.00")
    assert estado_resultados(desde=date(2026, 1, 1), hasta=HOY)["egresos"]["total"] == Decimal("0.00")
    assert _estado(carga, "3 Gastos") == ["crear"]


def test_movimiento_pagado_antes_del_arranque_es_error(centro, actor):
    carga = _nueva(_llenar(gastos=[{"fecha": date(2025, 12, 20), "descripcion": "Viejo", "monto": 100}]), actor)
    assert _estado(carga, "3 Gastos") == ["error"]


# ── Cobros ───────────────────────────────────────────────────────────────

def test_el_cobro_se_liga_solo_y_el_ingreso_no_se_cuenta_dos_veces(centro, actor, cliente_factory):
    from apps.contaduria.reportes import estado_resultados
    from apps.facturacion.models import Factura

    cliente_factory(razon_social="Heladería La Fuente")
    carga = _aplicar(_llenar(
        facturas=[{"folio": "F120", "cliente": "heladeria la fuente", "concepto": "Vasos",
                   "fecha": date(2026, 3, 1), "total": 8000, "regimen": "Exento"}],
        ingresos=[{"fecha": date(2026, 3, 15), "descripcion": "Depósito", "monto": 8000,
                   "cliente": "Heladería La Fuente"}],
    ), actor)
    fac = Factura.objects.get(folio_numero=120)
    assert fac.estado == "cobrada_total" and fac.cobros.count() == 1
    assert _saldo("cxc") == Decimal("0.00")
    assert estado_resultados(desde=date(2026, 1, 1), hasta=HOY)["ingresos"]["total"] == Decimal("8000.00")
    paso = next(p for p in carga.resumen["pasos"] if p["hoja"] == "2 Ingresos")
    assert "F120" in paso["ref"]


def test_cobro_ambiguo_pide_decidir(centro, actor, cliente_factory):
    cliente_factory(razon_social="Cliente A")
    carga = _nueva(_llenar(
        facturas=[{"folio": "F1", "cliente": "Cliente A", "concepto": "x", "fecha": date(2026, 3, 1), "total": 500, "regimen": "Exento"},
                  {"folio": "F2", "cliente": "Cliente A", "concepto": "y", "fecha": date(2026, 3, 2), "total": 500, "regimen": "Exento"}],
        ingresos=[{"fecha": date(2026, 3, 20), "descripcion": "Depósito", "monto": 500, "cliente": "Cliente A"}],
    ), actor)
    assert _estado(carga, "2 Ingresos") == ["error"]


def test_folio_existente_se_reconoce_y_recibe_su_cobro(centro, actor, cliente_factory):
    from apps.facturacion.models import Factura
    from apps.facturacion.services import fijar_total_con_impuestos

    cli = cliente_factory(razon_social="Restaurante Sol")
    fac = Factura.objects.create(cliente=cli, concepto="Menús", fecha_emision=date(2026, 4, 1),
                                 regimen_fiscal="exento", folio_numero=150)
    fijar_total_con_impuestos(fac, Decimal("3000"))
    Factura.objects.filter(pk=fac.pk).update(estado="emitida")
    carga = _aplicar(_llenar(
        facturas=[{"folio": "150", "cliente": "Restaurante Sol", "concepto": "Menús",
                   "fecha": date(2026, 4, 1), "total": 3000}],
        ingresos=[{"fecha": date(2026, 4, 10), "descripcion": "Pago", "monto": 1000, "factura": "F150"}],
    ), actor)
    assert _estado(carga, "4 Facturas") == ["saltar"]
    fac.refresh_from_db()
    assert fac.estado == "cobrada_parcial" and fac.monto_cobrado == Decimal("1000.00")


# ── Errores, pólizas, correos ────────────────────────────────────────────

def test_con_un_error_no_se_aplica_nada(centro, actor):
    from apps.contaduria.carga import motor
    from apps.tesoreria.models import Ingreso

    carga = _nueva(_llenar(ingresos=[
        {"fecha": date(2026, 2, 1), "descripcion": "Bien", "monto": 100},
        {"fecha": date(2026, 2, 1), "descripcion": "Mal", "monto": -5},
    ]), actor)
    assert carga.resumen["errores"] == 1
    with pytest.raises(motor.CargaConErrores):
        motor.aplicar(carga, actor)
    assert Ingreso.objects.count() == 0
    carga.refresh_from_db()
    assert carga.estado == "borrador"


def test_poliza_cuadra_o_no_entra(centro, actor):
    from apps.contaduria.models import Asiento

    buena = [{"poliza": "P1", "fecha": date(2026, 5, 1), "cuenta": "2.2.01", "cargo": 1600, "concepto": "Pago de IVA"},
             {"poliza": "P1", "fecha": date(2026, 5, 1), "cuenta": "1.1.02 · Bancos", "abono": 1600}]
    mala = [{"poliza": "P2", "fecha": date(2026, 5, 2), "cuenta": "2.2.01", "cargo": 100},
            {"poliza": "P2", "fecha": date(2026, 5, 2), "cuenta": "1.1.02", "abono": 90}]
    carga = _nueva(_llenar(polizas=buena + mala), actor)
    assert _estado(carga, "5 Pólizas") == ["crear", "error"]
    _aplicar(_llenar(polizas=buena), actor)
    assert Asiento.vigentes.filter(origen="carga").count() == 1
    assert _saldo("banco") == Decimal("-1600.00")


def test_no_manda_correos_a_clientes(centro, actor, monkeypatch):
    from django.db import transaction as _tx

    import lib.correos_auto as correos
    llamados = []
    monkeypatch.setattr(_tx, "on_commit", lambda fn, using=None, robust=False: fn())
    monkeypatch.setattr(correos, "enviar_confirmacion_pago", lambda *a, **k: llamados.append("pago"))
    monkeypatch.setattr(correos, "enviar_bienvenida", lambda *a, **k: llamados.append("bienvenida"))
    _aplicar(_llenar(
        facturas=[{"folio": "F7", "cliente": "Cliente Nuevo", "concepto": "x", "fecha": date(2026, 3, 1), "total": 700, "regimen": "Exento"}],
        ingresos=[{"fecha": date(2026, 3, 5), "descripcion": "Pago", "monto": 700, "factura": "F7"}],
    ), actor)
    assert llamados == []


# ── Deshacer ─────────────────────────────────────────────────────────────

def test_deshacer_deja_todo_como_estaba(centro, actor):
    from apps.contaduria.carga import motor
    from apps.facturacion.models import Factura

    contenido = _llenar(
        arranque={"1.1.02": 10000},
        facturas=[{"folio": "F55", "cliente": "Cliente Uno", "concepto": "x", "fecha": date(2026, 3, 1), "total": 1000, "regimen": "Exento"}],
        ingresos=[{"fecha": date(2026, 3, 5), "descripcion": "Pago", "monto": 400, "factura": "F55"}],
        gastos=[{"fecha": date(2026, 3, 6), "descripcion": "Tinta", "monto": 300}],
        hoy={"1.1.02": 10000},
    )
    carga = _aplicar(contenido, actor)
    motor.deshacer(carga, actor, "Me equivoqué de archivo")
    carga.refresh_from_db()
    assert carga.estado == "deshecha"
    for slot in ("banco", "cxc", "3.2.01", "ingreso_ventas", "egreso_operativo"):
        assert _saldo(slot) == Decimal("0.00"), slot
    assert Factura.objects.get(pk=carga.creados["facturas"][0]).folio_numero is None  # folio liberado
    # Y se puede volver a cargar el mismo archivo.
    otra = _aplicar(contenido, actor)
    assert Factura.objects.filter(folio_numero=55, estado="cobrada_parcial").exists()
    assert _estado(otra, "4 Facturas") == ["crear"]


def test_deshacer_se_niega_si_alguien_ya_cobro_despues(centro, actor):
    from apps.contaduria.carga import motor
    from apps.facturacion.models import Factura
    from apps.facturacion.services import registrar_cobro

    carga = _aplicar(_llenar(facturas=[{"folio": "F66", "cliente": "Cliente Dos", "concepto": "x",
                                        "fecha": date(2026, 3, 1), "total": 1000, "regimen": "Exento"}]), actor)
    registrar_cobro(Factura.objects.get(folio_numero=66), monto=Decimal("200"), fecha=date(2026, 9, 1),
                    metodo="transferencia", actor=actor)
    with pytest.raises(motor.CargaNoDeshacible):
        motor.deshacer(carga, actor, "prueba")


# ── Vistas y permisos ────────────────────────────────────────────────────

def test_vistas_de_punta_a_punta(client, centro, actor):
    from apps.contaduria.models import CargaContable
    from django.core.files.uploadedfile import SimpleUploadedFile

    client.force_login(actor)
    assert client.get("/contaduria/carga/").status_code == 200
    resp = client.get("/contaduria/carga/plantilla/")
    assert resp.status_code == 200 and "spreadsheetml" in resp["Content-Type"]
    contenido = _llenar(arranque={"1.1.02": 500})
    csv = "Fecha,Descripción,Cargo,Abono,Saldo\n05/02/2026,SPEI recibido,,300.00,800.00\n".encode()
    resp = client.post("/contaduria/carga/", {
        "plantilla": SimpleUploadedFile("mi.xlsx", contenido),
        "estado_1": SimpleUploadedFile("banco.csv", csv), "cuenta_1": str(_banco().pk),
        "crear_desde_estados": "1",
    })
    carga = CargaContable.objects.get()
    assert resp.status_code == 302 and resp["Location"].endswith(f"/contaduria/carga/{carga.pk}/")
    assert client.get(f"/contaduria/carga/{carga.pk}/").status_code == 200
    assert client.post(f"/contaduria/carga/{carga.pk}/aplicar/").status_code == 302
    carga.refresh_from_db()
    assert carga.estado == "aplicada"
    assert client.get(f"/contaduria/carga/{carga.pk}/").status_code == 200
    assert client.get(f"/contaduria/carga/{carga.pk}/archivo/").content == contenido
    estado = carga.estados_cuenta.get()
    assert client.get(f"/contaduria/carga/{carga.pk}/estado/{estado.pk}/").content == csv
    assert _saldo("banco") == Decimal("800.00")


def test_sin_permiso_no_entra(client, usuario_factory):
    disenador = usuario_factory(rol="disenador")
    client.force_login(disenador)
    assert client.get("/contaduria/carga/").status_code in (302, 403)
    assert client.get("/contaduria/carga/plantilla/").status_code in (302, 403)


def test_cargar_esta_en_los_defaults_de_quien_captura():
    from lib.permisos_defaults import DEFAULTS_POR_ROL

    for rol in ("super_admin", "dueno", "contador"):
        assert "cargar" in DEFAULTS_POR_ROL[rol]["contaduria"], rol
    assert "cargar" not in (DEFAULTS_POR_ROL.get("disenador", {}).get("contaduria") or [])


# ── Estados de cuenta ────────────────────────────────────────────────────

ESTADO_BBVA = """BBVA México
Cuenta: 0123456789
Periodo: 01/02/2026 - 28/02/2026

Fecha,Descripción,Referencia,Cargo,Abono,Saldo
01/02/2026,SPEI RECIBIDO HELADERIA,111,,5000.00,15000.00
03/02/2026,PAGO TINTA,222,1000.00,,14000.00
10/FEB/2026,COMISION MANEJO CUENTA,333,250.00,,13750.00
,TOTAL,,1250.00,5000.00,
"""


def test_estado_de_cuenta_se_casa_con_la_plantilla_y_completa_lo_que_falta(centro, actor):
    from apps.tesoreria.models import Egreso

    contenido = _llenar(
        arranque={"1.1.02": 10000},
        ingresos=[{"fecha": date(2026, 2, 1), "descripcion": "Pago de Heladería", "monto": 5000}],
        gastos=[{"fecha": date(2026, 2, 3), "descripcion": "Tinta", "monto": 1000}],
    )
    estados = [{"contenido": ESTADO_BBVA.encode("latin-1"), "nombre": "bbva-feb.csv", "cuenta": _banco()}]
    carga = _aplicar(contenido, actor, estados)
    movs = [p for p in carga.resumen["pasos"] if p["hoja"] == "Estados de cuenta"]
    assert [p["estado"] for p in movs] == ["saltar", "saltar", "crear"]   # 2 ya en la plantilla, 1 nueva
    comision = Egreso.vigentes.get(monto=Decimal("250.00"))
    assert "COMISION" in comision.descripcion
    # El saldo final del estado de cuenta cuadra solo: nada que ajustar.
    assert _saldo("banco") == Decimal("13750.00")
    assert carga.resumen["cuadre"]["lineas"][0]["diferencia"] == "0.00"


def test_estado_de_cuenta_sin_plantilla_y_sin_crear(centro, actor):
    from apps.tesoreria.models import Ingreso

    Ingreso.objects.create(monto=Decimal("5000"), fecha=date(2026, 2, 2), descripcion="Ya capturado")
    estados = [{"contenido": ESTADO_BBVA.encode(), "nombre": "bbva.csv", "cuenta": _banco()}]
    carga = _nueva(None, actor, estados, crear=False)
    movs = [p for p in carga.resumen["pasos"] if p["hoja"] == "Estados de cuenta"]
    assert [p["estado"] for p in movs] == ["saltar", "saltar", "saltar"]
    assert "ING-" in movs[0]["mensajes"][0]
    # Sin crear, la diferencia contra el saldo final queda como ajuste visible.
    assert carga.resumen["cuadre"]["lineas"][0]["declarado"] == "13750.00"


def test_estado_de_cuenta_en_excel_y_en_pdf(centro, actor):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Banorte"])
    ws.append(["FECHA", "DESCRIPCIÓN", "DEPÓSITOS", "RETIROS", "SALDO"])
    ws.append([date(2026, 3, 2), "DEPOSITO EFECTIVO", 700, None, 700])
    buf = io.BytesIO()
    wb.save(buf)
    estados = [{"contenido": buf.getvalue(), "nombre": "banorte.xlsx", "cuenta": _banco()},
               {"contenido": b"%PDF-1.4 ...", "nombre": "marzo.pdf", "cuenta": _banco()}]
    carga = _aplicar(None, actor, estados)
    assert _saldo("banco") == Decimal("700.00")
    assert any("PDF" in a for a in carga.resumen["avisos"])
    assert carga.estados_cuenta.count() == 2   # el PDF queda como evidencia


def test_estados_que_se_enciman_no_duplican(centro, actor):
    estados = [{"contenido": ESTADO_BBVA.encode(), "nombre": "a.csv", "cuenta": _banco()},
               {"contenido": ESTADO_BBVA.encode(), "nombre": "b.csv", "cuenta": _banco()}]
    carga = _aplicar(None, actor, estados)
    creados = [p for p in carga.resumen["pasos"] if p["hoja"] == "Estados de cuenta" and p["estado"] == "crear"]
    assert len(creados) == 3
    assert _saldo("banco") == Decimal("13750.00")


def test_lectura_de_fechas_bancarias():
    from apps.contaduria.carga import lectura
    assert lectura.leer_fecha("10/FEB/2026") == date(2026, 2, 10)
    assert lectura.leer_fecha("02-ene-26") == date(2026, 1, 2)


def test_el_chalan_sabe_como_quedo_la_carga(centro, actor, usuario_factory):
    from capacidades.gating import gate_ok
    from capacidades.lecturas import _LECTURAS

    cap = _LECTURAS["contaduria_carga"]
    assert cap.fn({}, actor)["cargas"] == 0
    _aplicar(_llenar(arranque={"1.1.02": 100}, hoy={"1.1.02": 100}), actor)
    r = cap.fn({}, actor)
    assert r["estado"] == "Aplicada" and r["errores"] == 0
    assert any(d["cuenta"].startswith("1.1.02") for d in r["diferencias"])
    assert gate_ok(cap.gating, actor)
    assert not gate_ok(cap.gating, usuario_factory(rol="disenador"))


# ── El Chalán revisa lo que se sube ──────────────────────────────────────

ESTADO_CON_RAROS = """Fecha,Descripción,Cargo,Abono,Saldo
05/03/2026,TRASPASO A CAJA CHICA,2000.00,,8000.00
06/03/2026,SPEI RECIBIDO RESTAURANTE SOL,,3000.00,11000.00
07/03/2026,PAGO SAT IVA FEBRERO,1600.00,,9400.00
08/03/2026,COMPRA PAPELERIA,500.00,,8900.00
"""


def test_el_chalan_clasifica_lo_que_solo_trae_el_banco(centro, actor, monkeypatch, cliente_factory):
    from apps.tesoreria.models import CentroDeCosto, Egreso, Ingreso

    CentroDeCosto.objects.create(nombre="Insumos", slug="insumos", naturaleza="proyecto")
    cliente_factory(razon_social="Restaurante Sol")
    preguntas = _chalan_dice(monkeypatch, {
        "CAJA CHICA": {"tipo": "traspaso", "cuenta": "1.1.01", "confianza": 0.95},
        "RESTAURANTE SOL": {"tipo": "cobro", "factura": "F300", "cliente": "Restaurante Sol", "confianza": 0.9},
        "PAGO SAT": {"tipo": "impuesto", "cuenta": "2.2.01", "concepto": "Pago de IVA", "confianza": 0.9},
        "PAPELERIA": {"tipo": "gasto", "centro": "insumos", "confianza": 0.8},
    })
    contenido = _llenar(arranque={"1.1.02": 10000}, facturas=[
        {"folio": "F300", "cliente": "Restaurante Sol", "concepto": "Menús", "fecha": date(2026, 3, 1),
         "total": 3000, "regimen": "Exento"}])
    estados = [{"contenido": ESTADO_CON_RAROS.encode(), "nombre": "marzo.csv", "cuenta": _banco()}]
    carga = _aplicar(contenido, actor, estados)
    assert preguntas, "El Chalán debió revisar los movimientos"
    assert _saldo("caja") == Decimal("2000.00")                      # traspaso, no gasto
    assert _saldo("iva_trasladado") == Decimal("-1600.00")            # pago de IVA, no gasto
    assert Egreso.vigentes.count() == 1                               # sólo la papelería
    assert Egreso.vigentes.get().centro_de_costo.slug == "insumos"
    cobro = Ingreso.vigentes.get(monto=Decimal("3000"))
    assert cobro.factura.folio_numero == 300 and cobro.factura.estado == "cobrada_total"
    assert _saldo("banco") == Decimal("8900.00")
    assert carga.resumen["ia"]["revisados"] == 4
    marcados = [m for p in carga.resumen["pasos"] for m in p["mensajes"] if m.startswith("🤖")]
    assert len(marcados) == 4


def test_lo_que_el_chalan_inventa_se_descarta(centro, actor, monkeypatch):
    from apps.tesoreria.models import Egreso

    _chalan_dice(monkeypatch, {
        "CAJA CHICA": {"tipo": "traspaso", "cuenta": "9.9.99", "confianza": 0.99},     # cuenta que no existe
        "RESTAURANTE": {"tipo": "cobro", "factura": "F999", "confianza": 0.99},         # factura que no existe
        "PAGO SAT": {"tipo": "impuesto", "cuenta": "2.2.01", "confianza": 0.2},        # no está seguro
        "PAPELERIA": {"tipo": "gasto", "centro": "inventado", "confianza": 0.9},      # centro que no existe
    })
    estados = [{"contenido": ESTADO_CON_RAROS.encode(), "nombre": "marzo.csv", "cuenta": _banco()}]
    carga = _aplicar(None, actor, estados)
    # Todo cae a la regla de siempre: nada se inventa.
    assert Egreso.vigentes.count() == 3
    # Al comodín «Otros», no al primero del alfabeto («Impuestos y comisiones»).
    assert set(Egreso.vigentes.values_list("centro_de_costo__slug", flat=True)) == {"otros"}
    assert _saldo("caja") == Decimal("0.00")
    mensajes = " ".join(m for p in carga.resumen["pasos"] for m in p["mensajes"])
    assert "no estuvo seguro" in mensajes and "no con datos del catálogo" in mensajes


def test_aplicar_usa_lo_que_el_chalan_dijo_en_la_vista_previa(centro, actor, monkeypatch):
    """La IA se pregunta al subir; aplicar no le vuelve a preguntar (no puede
    cambiar de opinión entre lo que se vio y lo que se guarda)."""
    from apps.contaduria.carga import ia, motor

    _chalan_dice(monkeypatch, {"CAJA CHICA": {"tipo": "traspaso", "cuenta": "1.1.01", "confianza": 0.95}})
    estados = [{"contenido": ESTADO_CON_RAROS.encode(), "nombre": "marzo.csv", "cuenta": _banco()}]
    carga = _nueva(None, actor, estados)

    def _prohibido(prompt, usuario):
        raise AssertionError("aplicar no debe llamar a la IA")
    monkeypatch.setattr(ia, "_llamar", _prohibido)
    motor.aplicar(carga, actor)
    assert _saldo("caja") == Decimal("2000.00")


def test_el_chalan_sugiere_centro_y_desempata_cobros(centro, actor, monkeypatch, cliente_factory):
    from apps.facturacion.models import Factura
    from apps.tesoreria.models import CentroDeCosto, Egreso

    CentroDeCosto.objects.create(nombre="Insumos", slug="insumos", naturaleza="proyecto")
    cliente_factory(razon_social="Cliente A")
    _chalan_dice(monkeypatch, {
        "Tinta para plotter": {"tipo": "gasto", "centro": "insumos", "confianza": 0.9},
        "Pago factura dos": {"tipo": "cobro", "factura": "F2", "confianza": 0.85},
    })
    carga = _aplicar(_llenar(
        facturas=[{"folio": "F1", "cliente": "Cliente A", "concepto": "x", "fecha": date(2026, 3, 1), "total": 500, "regimen": "Exento"},
                  {"folio": "F2", "cliente": "Cliente A", "concepto": "y", "fecha": date(2026, 3, 2), "total": 500, "regimen": "Exento"}],
        ingresos=[{"fecha": date(2026, 3, 20), "descripcion": "Pago factura dos", "monto": 500, "cliente": "Cliente A"}],
        gastos=[{"fecha": date(2026, 3, 4), "descripcion": "Tinta para plotter", "monto": 900}],
    ), actor)
    assert Egreso.vigentes.get().centro_de_costo.slug == "insumos"
    assert Factura.objects.get(folio_numero=2).estado == "cobrada_total"
    assert Factura.objects.get(folio_numero=1).estado == "emitida"
    assert carga.resumen["errores"] == 0
