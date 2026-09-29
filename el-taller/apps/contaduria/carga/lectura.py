"""Lee la plantilla llena.

Tolerante con lo que pega una persona: fechas «3/2/26», «2026-02-03» o de
Excel; montos «$1,200.50» o «(300.00)»; títulos con o sin acento o «*»; filas
vacías intercaladas. Lo que no entiende no lo adivina: lo devuelve como error
del renglón, en español, para que se corrija en el Excel.

No toca la base: devuelve estructuras planas que el motor resuelve.
"""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from . import esquema as E

MAX_FILAS = 5000


class PlantillaInvalida(ValueError):
    """El archivo no es la plantilla (o no se puede abrir)."""


@dataclass
class Fila:
    hoja: str
    numero: int                    # número de renglón en Excel (para decir «fila 12»)
    valores: dict = field(default_factory=dict)
    errores: list[str] = field(default_factory=list)


@dataclass
class HojaSaldos:
    fecha: date | None = None
    error_fecha: str = ""
    filas: list[Fila] = field(default_factory=list)


@dataclass
class Libro:
    arranque: HojaSaldos = field(default_factory=HojaSaldos)
    hoy: HojaSaldos = field(default_factory=HojaSaldos)
    renglones: dict[str, list[Fila]] = field(default_factory=dict)
    hojas_faltantes: list[str] = field(default_factory=list)


# ── Normalización ────────────────────────────────────────────────────────

def normalizar(texto) -> str:
    """«¿Incluye IVA? *» → «incluye iva». Sin acentos, minúsculas, sin signos."""
    s = unicodedata.normalize("NFKD", str(texto or ""))
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower()
    s = re.sub(r"[^a-z0-9ñ@.\- ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_FORMATOS_FECHA = ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d-%m-%y", "%d.%m.%Y", "%Y/%m/%d")
_EPOCA_EXCEL = date(1899, 12, 30)
_MESES = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
          "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12}


def leer_fecha(valor) -> date | None:
    """`None` si viene vacío. Lanza ValueError con mensaje si no se entiende."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, int | float) and 20000 < valor < 80000:
        return _EPOCA_EXCEL + timedelta(days=int(valor))
    texto = str(valor).strip()
    for fmt in _FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    # Los bancos escriben «02/ENE/2026», «02-ene-26» o «2 enero 2026».
    m = re.fullmatch(r"(\d{1,2})[\s/.\-]+([a-zA-Záéíóú]{3,10})\.?[\s/.\-]+(\d{2,4})", texto)
    if m and normalizar(m.group(2))[:3] in _MESES:
        anio = int(m.group(3))
        anio += 2000 if anio < 100 else 0
        try:
            return date(anio, _MESES[normalizar(m.group(2))[:3]], int(m.group(1)))
        except ValueError:
            pass
    raise ValueError(f"no entiendo la fecha «{texto}» (usa 31/01/2026)")


def leer_monto(valor) -> Decimal | None:
    """`None` si viene vacío. Acepta 1200, 1,200.50, $1,200.50, (300.00) = −300."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        raise ValueError("un monto no puede ser Sí/No")
    if isinstance(valor, int | float | Decimal):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    texto = str(valor).strip().replace("$", "").replace("MXN", "").replace(" ", "")
    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()").replace(",", "")
    try:
        monto = Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise ValueError(f"no entiendo el monto «{valor}»") from None
    return -monto if negativo else monto


_SI = {"si", "s", "x", "1", "true", "verdadero", "yes", "y"}
_NO = {"no", "n", "0", "false", "falso", ""}


def leer_si_no(valor) -> bool:
    if isinstance(valor, bool):
        return valor
    n = normalizar(valor)
    if n in _SI:
        return True
    if n in _NO:
        return False
    raise ValueError(f"escribe Sí o No (vino «{valor}»)")


def leer_texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# ── Libro ────────────────────────────────────────────────────────────────

def leer(contenido: bytes) -> Libro:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 — cualquier cosa que no abra es «no es la plantilla»
        raise PlantillaInvalida(
            "No pude abrir el archivo. Sube la plantilla en formato Excel (.xlsx)."
        ) from exc

    por_nombre = {normalizar(n): n for n in wb.sheetnames}

    def hoja(nombre: str):
        real = por_nombre.get(normalizar(nombre))
        return wb[real] if real else None

    reconocidas = [n for n in (E.HOJA_ARRANQUE, E.HOJA_HOY, *E.HOJAS_RENGLONES) if hoja(n) is not None]
    if not reconocidas:
        raise PlantillaInvalida(
            "El archivo no tiene ninguna de las hojas de la plantilla («1 Arranque», «2 Ingresos»…). "
            "Descarga la plantilla desde esta pantalla y llénala."
        )

    libro = Libro()
    for nombre in (E.HOJA_ARRANQUE, E.HOJA_HOY):
        ws = hoja(nombre)
        if ws is None:
            libro.hojas_faltantes.append(nombre)
            continue
        destino = _leer_saldos(ws, nombre)
        if nombre == E.HOJA_ARRANQUE:
            libro.arranque = destino
        else:
            libro.hoy = destino
    for nombre, columnas in E.HOJAS_RENGLONES.items():
        ws = hoja(nombre)
        if ws is None:
            libro.hojas_faltantes.append(nombre)
            libro.renglones[nombre] = []
            continue
        libro.renglones[nombre] = _leer_renglones(ws, nombre, columnas)
    wb.close()
    return libro


def _mapa_columnas(encabezado: list, columnas) -> dict[str, int]:
    """{clave: índice} por título o alias normalizado. El título exacto gana
    sobre un alias (en Gastos «Cargo» es alias de Monto; en Pólizas es columna)."""
    normales = [normalizar(v) for v in encabezado]
    mapa: dict[str, int] = {}
    for col in columnas:
        titulo = normalizar(col.titulo)
        if titulo in normales:
            mapa[col.clave] = normales.index(titulo)
    usados = set(mapa.values())
    for col in columnas:
        if col.clave in mapa:
            continue
        for alias in col.alias:
            a = normalizar(alias)
            if a in normales and normales.index(a) not in usados:
                mapa[col.clave] = normales.index(a)
                usados.add(mapa[col.clave])
                break
    return mapa


def _convertir(col: E.Columna, crudo, fila: Fila) -> object:
    try:
        if col.tipo == "fecha":
            return leer_fecha(crudo)
        if col.tipo == "monto":
            return leer_monto(crudo)
        if col.tipo == "si_no":
            return leer_si_no(crudo)
        return leer_texto(crudo)
    except ValueError as exc:
        fila.errores.append(f"«{col.titulo}»: {exc}")
        return None


def _vacia(valores) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in valores)


def _leer_renglones(ws, nombre: str, columnas) -> list[Fila]:
    filas = list(ws.iter_rows(values_only=True))
    # El encabezado es la primera fila (entre las 10 primeras) que trae el
    # título de la primera columna obligatoria: resiste que alguien borre el
    # título o la ayuda de arriba.
    ancla = normalizar(next(c.titulo for c in columnas if c.requerido))
    idx_enc = next(
        (i for i, f in enumerate(filas[:10]) if f and ancla in [normalizar(v) for v in f]),
        None,
    )
    if idx_enc is None:
        return [Fila(nombre, 1, errores=[f"No encuentro los títulos de columna (busqué «{ancla}»)."])]
    mapa = _mapa_columnas(list(filas[idx_enc]), columnas)
    faltan = [c.titulo for c in columnas if c.requerido and c.clave not in mapa]
    if faltan:
        return [Fila(nombre, idx_enc + 1, errores=[f"Faltan las columnas: {', '.join(faltan)}."])]

    salida: list[Fila] = []
    for i, crudo in enumerate(filas[idx_enc + 1:], start=idx_enc + 2):
        if crudo is None or _vacia(crudo):
            continue
        if len(salida) >= MAX_FILAS:
            salida.append(Fila(nombre, i, errores=[f"La hoja pasa de {MAX_FILAS} renglones: pártela en dos cargas."]))
            break
        fila = Fila(nombre, i)
        for col in columnas:
            j = mapa.get(col.clave)
            valor = crudo[j] if j is not None and j < len(crudo) else None
            fila.valores[col.clave] = _convertir(col, valor, fila)
        for col in columnas:
            if col.requerido and fila.valores.get(col.clave) in (None, "") and not any(
                col.titulo in e for e in fila.errores
            ):
                fila.errores.append(f"falta «{col.titulo}».")
        salida.append(fila)
    return salida


def _leer_saldos(ws, nombre: str) -> HojaSaldos:
    res = HojaSaldos()
    try:
        res.fecha = leer_fecha(ws[E.CELDA_FECHA].value)
    except ValueError as exc:
        res.error_fecha = str(exc)
    except Exception:  # noqa: BLE001 — hoja de sólo lectura sin esa celda
        res.fecha = None
    filas = list(ws.iter_rows(values_only=True))
    idx_enc = next(
        (i for i, f in enumerate(filas[:10]) if f and normalizar(f[0]) in {"cuenta", "cuenta *"}),
        None,
    )
    if idx_enc is None:
        return res
    mapa = _mapa_columnas(list(filas[idx_enc]), E.COLUMNAS_SALDOS)
    for i, crudo in enumerate(filas[idx_enc + 1:], start=idx_enc + 2):
        if crudo is None or _vacia(crudo):
            continue
        fila = Fila(nombre, i)
        for col in E.COLUMNAS_SALDOS:
            j = mapa.get(col.clave)
            valor = crudo[j] if j is not None and j < len(crudo) else None
            fila.valores[col.clave] = _convertir(col, valor, fila)
        # Una cuenta precargada sin saldo es «no lo sé»: no se toca. Sólo
        # cuenta el renglón si trae saldo (0 incluido: 0 es cero).
        if fila.valores.get("saldo") is None and not fila.errores:
            continue
        res.filas.append(fila)
    return res
