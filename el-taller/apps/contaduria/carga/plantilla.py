"""Genera la plantilla Excel de la carga contable.

Se arma con el catálogo VIVO (cuentas, centros de costo, clientes, proveedores),
así las listas desplegables ofrecen exactamente lo que El Despacho reconoce.
Las listas de clientes y proveedores sugieren pero no obligan: un cliente que
no existe se da de alta al cargar sus facturas.
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import esquema as E

FILAS_CON_FORMATO = 1500
FORMATO_FECHA = "dd/mm/yyyy"
FORMATO_MONTO = "#,##0.00;[Red]-#,##0.00"

_FUENTE_TITULO = Font(bold=True, size=14)
_FUENTE_AYUDA = Font(italic=True, size=9, color="667085")
_FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="465FFF")
_RELLENO_REQUERIDO = PatternFill("solid", fgColor="2A31D8")
_RELLENO_FECHA = PatternFill("solid", fgColor="FEF0C7")


def etiqueta_cuenta(cuenta) -> str:
    """Cómo se escribe una cuenta en la plantilla: «1.1.02 · Bancos»."""
    return f"{cuenta.codigo} · {cuenta.nombre}"


def generar() -> bytes:
    """El .xlsx listo para descargar."""
    listas = _valores_de_listas()
    wb = Workbook()
    _hoja_leeme(wb.active)

    hoja_listas = wb.create_sheet(E.HOJA_LISTAS)
    rangos = _escribir_listas(hoja_listas, listas)
    hoja_listas.sheet_state = "hidden"

    from apps.contaduria.models import CuentaContable

    arranque = [
        c for c in CuentaContable.activas.filter(tipo__in=["activo", "pasivo", "capital"]).order_by("codigo")
        if c.slot not in E.SLOTS_DEL_DETALLE and c.slot != "ajuste_captura"
        and c.codigo not in (E.CODIGO_UTILIDADES_ACUMULADAS, "3.2.02")
    ]
    hoy = list(CuentaContable.activas.filter(tipo="activo", codigo__startswith="1.1").order_by("codigo"))

    _hoja_saldos(
        wb, E.HOJA_ARRANQUE, rangos,
        titulo="1 · Saldos de arranque",
        ayuda_fecha="Fecha de arranque: el primer día desde el que El Despacho lleva TODO (recomendado: 1 de enero).",
        ayuda="Lo que había en cada cuenta al EMPEZAR ese día. Vacío = no tocar · 0 = vale cero. "
              "Clientes, proveedores y reembolsos NO van aquí: salen de las hojas 3 y 4.",
        fecha=date(date.today().year, 1, 1),
        cuentas=arranque,
    )
    for nombre, columnas in E.HOJAS_RENGLONES.items():
        _hoja_renglones(wb, nombre, columnas, rangos)
    _hoja_saldos(
        wb, E.HOJA_HOY, rangos,
        titulo="6 · Saldos reales hoy",
        ayuda_fecha="Fecha de la comprobación (normalmente hoy).",
        ayuda="El saldo REAL de cada cuenta en esa fecha (estado de cuenta, arqueo de caja). "
              "Lo que falte para cuadrar se registra como ajuste; lo verás antes de aplicar.",
        fecha=date.today(),
        cuentas=hoy,
    )
    # Listas al final de las pestañas (está oculta, pero el orden importa al abrirla).
    wb.move_sheet(E.HOJA_LISTAS, offset=len(wb.sheetnames))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Listas ───────────────────────────────────────────────────────────────

def _valores_de_listas() -> dict[str, list[str]]:
    from apps.contaduria.models import CuentaContable
    from apps.el_catalogo.models import Proveedor
    from apps.la_cartera.models import Cliente
    from apps.tesoreria.models import CentroDeCosto
    from apps.tesoreria.models.egreso import METODOS_EGRESO
    from apps.tesoreria.models.ingreso import METODOS_INGRESO

    return {
        "si_no": [e for _, e in E.SI_NO],
        "metodos_ingreso": [e for _, e in METODOS_INGRESO],
        "metodos_gasto": [e for _, e in METODOS_EGRESO],
        "estados_gasto": [e for _, e in E.ESTADOS_GASTO],
        "regimenes": [e for _, e in E.REGIMENES_FACTURA],
        "centros": list(CentroDeCosto.objects.filter(activo=True).order_by("nombre")
                        .values_list("nombre", flat=True)),
        "cuentas": [etiqueta_cuenta(c) for c in CuentaContable.activas.order_by("codigo")],
        "clientes": list(Cliente.objects.filter(activo=True).order_by("razon_social").values_list("razon_social", flat=True)),
        "proveedores": list(Proveedor.objects.order_by("razon_social").values_list("razon_social", flat=True)),
    }


def _escribir_listas(ws, listas: dict[str, list[str]]) -> dict[str, str]:
    """Una columna por lista; devuelve {nombre: rango absoluto} para validar."""
    rangos: dict[str, str] = {}
    for i, (nombre, valores) in enumerate(listas.items(), start=1):
        letra = get_column_letter(i)
        ws.cell(row=1, column=i, value=nombre).font = Font(bold=True)
        for j, v in enumerate(valores, start=2):
            ws.cell(row=j, column=i, value=v)
        ultima = max(2, len(valores) + 1)
        rangos[nombre] = f"'{E.HOJA_LISTAS}'!${letra}$2:${letra}${ultima}"
        ws.column_dimensions[letra].width = 30
    return rangos


def _validar(ws, col: E.Columna, letra: str, primera: int, rangos: dict[str, str]) -> None:
    nombre = col.lista or ("si_no" if col.tipo == "si_no" else "")
    if not nombre or nombre not in rangos:
        return
    dv = DataValidation(type="list", formula1=f"={rangos[nombre]}", allow_blank=True)
    dv.showErrorMessage = col.estricta
    if col.estricta:
        dv.errorTitle = "Valor no reconocido"
        dv.error = "Elige un valor de la lista."
    ws.add_data_validation(dv)
    dv.add(f"{letra}{primera}:{letra}{primera + FILAS_CON_FORMATO}")


# ── Hojas ────────────────────────────────────────────────────────────────

def _encabezado(ws, fila: int, columnas, rangos, primera_dato: int) -> None:
    for i, col in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        celda = ws.cell(row=fila, column=i, value=col.titulo + (" *" if col.requerido else ""))
        celda.font = _FUENTE_ENCABEZADO
        celda.fill = _RELLENO_REQUERIDO if col.requerido else _RELLENO_ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        if col.ayuda:
            celda.comment = Comment(col.ayuda, "El Despacho")
        ws.column_dimensions[letra].width = col.ancho
        formato = FORMATO_FECHA if col.tipo == "fecha" else FORMATO_MONTO if col.tipo == "monto" else None
        if formato:
            for r in range(primera_dato, primera_dato + FILAS_CON_FORMATO):
                ws.cell(row=r, column=i).number_format = formato
        _validar(ws, col, letra, primera_dato, rangos)
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def _hoja_renglones(wb, nombre: str, columnas, rangos) -> None:
    ws = wb.create_sheet(nombre)
    titulos = {
        E.HOJA_INGRESOS: ("2 · Ingresos", "Cada entrada de dinero desde la fecha de arranque. Lo que ya está en El Despacho se reconoce solo."),
        E.HOJA_GASTOS: ("3 · Gastos", "Cada gasto desde la fecha de arranque, más los que al arrancar seguían sin pagarse o por reembolsar."),
        E.HOJA_FACTURAS: ("4 · Facturas", "Las emitidas desde el arranque y las anteriores que seguían sin cobrarse. Sus cobros van en la hoja 2."),
        E.HOJA_POLIZAS: ("5 · Pólizas (opcional, para el contador)", "Lo que no es ingreso ni gasto: pago de impuestos, préstamos, aportaciones. Cada póliza debe cuadrar."),
    }
    titulo, ayuda = titulos[nombre]
    ws["A1"] = titulo
    ws["A1"].font = _FUENTE_TITULO
    ws["A2"] = ayuda + "  (* = obligatorio · pasa el mouse sobre cada título para ver su ayuda)"
    ws["A2"].font = _FUENTE_AYUDA
    _encabezado(ws, E.FILA_ENCABEZADO, columnas, rangos, E.FILA_ENCABEZADO + 1)


def _hoja_saldos(wb, nombre, rangos, *, titulo, ayuda_fecha, ayuda, fecha, cuentas) -> None:
    ws = wb.create_sheet(nombre)
    ws["A1"] = titulo
    ws["A1"].font = _FUENTE_TITULO
    ws["A2"] = "Fecha:"
    ws["A2"].font = Font(bold=True)
    ws[E.CELDA_FECHA] = fecha
    ws[E.CELDA_FECHA].number_format = FORMATO_FECHA
    ws[E.CELDA_FECHA].fill = _RELLENO_FECHA
    ws[E.CELDA_FECHA].comment = Comment(ayuda_fecha, "El Despacho")
    ws["C2"] = ayuda_fecha
    ws["C2"].font = _FUENTE_AYUDA
    ws["A3"] = ayuda
    ws["A3"].font = _FUENTE_AYUDA
    primera = E.FILA_ENCABEZADO_SALDOS + 1
    _encabezado(ws, E.FILA_ENCABEZADO_SALDOS, E.COLUMNAS_SALDOS, rangos, primera)
    for i, c in enumerate(cuentas):
        ws.cell(row=primera + i, column=1, value=etiqueta_cuenta(c))


def _hoja_leeme(ws) -> None:
    ws.title = E.HOJA_LEEME
    ws.column_dimensions["A"].width = 110
    lineas = [
        ("Carga contable de El Despacho", _FUENTE_TITULO),
        ("Para subir de un jalón la contabilidad que se llevó fuera y que, desde ahí, todo cuadre.", None),
        ("", None),
        ("Cómo se llena", Font(bold=True, size=12)),
        ("1. Hoja «1 Arranque»: escribe la FECHA DE ARRANQUE (el primer día desde el que El Despacho debe tenerlo "
         "todo; lo mejor es el 1 de enero, así el estado de resultados del año sale completo) y lo que había en "
         "cada cuenta al empezar ese día: bancos, caja, capital. Vacío = no tocar; 0 = vale cero.", None),
        ("2. Hojas «2 Ingresos» y «3 Gastos»: cada movimiento desde la fecha de arranque hasta hoy. Puedes pegar "
         "directo del estado de cuenta. Lo que ya está capturado en El Despacho se reconoce solo (mismo monto, "
         "±3 días) y NO se duplica; si de verdad son dos, escribe «Sí» en «Forzar».", None),
        ("   En Gastos, lo que al arrancar seguía sin pagarse (o por reembolsar a alguien) va con su fecha original "
         "y su estado: así queda como cuenta por pagar.", None),
        ("3. Hoja «4 Facturas»: las emitidas desde el arranque y las anteriores que seguían sin cobrarse (con lo que "
         "ya habían pagado en «Cobrado antes del arranque»). Los cobros de las nuevas van en la hoja 2, en «Factura "
         "que paga»; si lo dejas vacío y el cobro es inequívoco (mismo cliente y monto), se liga solo.", None),
        ("4. Hoja «5 Pólizas» (opcional, para el contador): lo que no es ingreso ni gasto — pago de impuestos, "
         "préstamos, aportaciones, depreciación. Cada póliza debe cuadrar (cargos = abonos).", None),
        ("5. Hoja «6 Saldos hoy»: el saldo REAL de bancos y caja hoy. Si después de todo lo anterior algo no cuadra, "
         "la diferencia se registra como ajuste. Te la enseñamos antes de aplicar.", None),
        ("", None),
        ("Qué pasa al subirlo", Font(bold=True, size=12)),
        ("· Primero ves una VISTA PREVIA exacta: qué se crea, qué se reconoce como repetido, qué tiene error y cómo "
         "queda el balance. Nada se guarda todavía.", None),
        ("· Si un renglón tiene error, corrígelo aquí y vuelve a subir el archivo. Con errores no se puede aplicar: "
         "una contabilidad a medias es justo la que no cuadra.", None),
        ("· Volver a subir el mismo archivo no duplica: cada renglón ya cargado se reconoce.", None),
        ("· Una carga aplicada se puede DESHACER completa, mientras nadie haya cobrado o pagado sobre lo importado.", None),
        ("· No manda correos a clientes (ni «recibimos tu pago» ni de bienvenida) y no toca lo que ya existe.", None),
        ("", None),
        ("Formatos", Font(bold=True, size=12)),
        ("· Fechas: 31/01/2026, 2026-01-31 o una fecha de Excel. Montos: 1200, 1,200.50 o $1,200.50.", None),
        ("· Clientes y proveedores: como están en El Despacho (o el RFC). La lista sugiere; un cliente nuevo de una "
         "factura se da de alta solo.", None),
        ("· No cambies el nombre de las pestañas ni de los títulos de columna.", None),
    ]
    for i, (texto, fuente) in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fuente is not None:
            c.font = fuente
