"""Lee un estado de cuenta tal como lo exporta la banca en línea.

CSV o Excel (.xlsx). Los bancos ponen arriba un preámbulo (nombre del
cliente, número de cuenta, periodo) y después la tabla; aquí se busca el
renglón de títulos entre los primeros 40 y se reconocen las columnas por
nombre: «Fecha», «Descripción/Concepto», y el importe como «Monto» firmado o
como el par «Depósitos/Abonos» – «Retiros/Cargos». Ojo: en el estado de cuenta
«Cargo» es dinero que SALE y «Abono» dinero que ENTRA (al revés que en el
libro del despacho, porque el banco lo ve desde su lado).

Un PDF no se puede leer con confianza: se guarda como evidencia de la carga y
se avisa que para usarlo hace falta la exportación en CSV o Excel.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from . import lectura

CERO = Decimal("0.00")

_FECHA = ("fecha", "fecha operacion", "fecha de operacion", "f operacion", "f. operacion", "fecha movimiento", "dia", "date")
_DESC = ("descripcion", "concepto", "detalle", "movimiento", "descripcion del movimiento", "description")
_MONTO = ("monto", "importe", "amount", "cantidad")
_DEPOSITO = ("deposito", "depositos", "abono", "abonos", "ingreso", "ingresos", "credito", "creditos")
_RETIRO = ("retiro", "retiros", "cargo", "cargos", "egreso", "egresos", "debito", "debitos")
_SALDO = ("saldo", "saldo final", "saldo disponible", "balance")
_REF = ("referencia", "folio", "ref", "clave de rastreo", "rastreo", "numero de referencia")


@dataclass
class Movimiento:
    fila: int
    fecha: date
    descripcion: str
    monto: Decimal                 # firmado: + entra a la cuenta, − sale
    saldo: Decimal | None = None
    referencia: str = ""


@dataclass
class EstadoCuenta:
    nombre: str
    movimientos: list[Movimiento] = field(default_factory=list)
    error: str = ""
    ignorados: int = 0

    @property
    def saldo_final(self) -> tuple[date, Decimal] | None:
        """El saldo del último día, respetando si el banco lista del más
        viejo al más nuevo o al revés."""
        con_saldo = [m for m in self.movimientos if m.saldo is not None]
        if not con_saldo:
            return None
        ultimo_dia = max(m.fecha for m in con_saldo)
        del_dia = [m for m in con_saldo if m.fecha == ultimo_dia]
        descendente = self.movimientos[0].fecha > self.movimientos[-1].fecha
        m = del_dia[0] if descendente else del_dia[-1]
        return m.fecha, m.saldo


def _coincide(titulo: str, candidatos) -> bool:
    return any(titulo == c or titulo.startswith(c + " ") for c in candidatos)


def _columna(titulos: list[str], candidatos, usadas: set[int]) -> int | None:
    for i, t in enumerate(titulos):
        if i not in usadas and _coincide(t, candidatos):
            usadas.add(i)
            return i
    return None


def _filas_de(contenido: bytes, nombre: str) -> list[list]:
    if nombre.lower().endswith(".xlsx") or contenido[:2] == b"PK":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        filas = [list(f) for f in ws.iter_rows(values_only=True)]
        wb.close()
        return filas
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4096]
    delim = max((";", ",", "\t", "|"), key=muestra.count)
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delim)]


def leer(contenido: bytes, nombre: str) -> EstadoCuenta:
    res = EstadoCuenta(nombre=nombre)
    bajo = (nombre or "").lower()
    if bajo.endswith(".pdf") or contenido[:4] == b"%PDF":
        res.error = ("Es un PDF: se guarda como evidencia, pero no se puede leer. Descarga el estado de cuenta "
                     "en CSV o Excel desde la banca en línea y súbelo así.")
        return res
    if bajo.endswith(".xls") or contenido[:4] == b"\xd0\xcf\x11\xe0":
        res.error = "Es un Excel antiguo (.xls): ábrelo y guárdalo como .xlsx o CSV."
        return res
    try:
        filas = _filas_de(contenido, nombre)
    except Exception:  # noqa: BLE001 — cualquier archivo ilegible
        res.error = "No pude abrir el archivo: súbelo en CSV o Excel (.xlsx)."
        return res

    idx = cols = None
    for i, fila in enumerate(filas[:40]):
        titulos = [lectura.normalizar(v) for v in fila]
        usadas: set[int] = set()
        c_fecha = _columna(titulos, _FECHA, usadas)
        if c_fecha is None:
            continue
        c = {"fecha": c_fecha, "monto": _columna(titulos, _MONTO, usadas),
             "deposito": _columna(titulos, _DEPOSITO, usadas), "retiro": _columna(titulos, _RETIRO, usadas),
             "saldo": _columna(titulos, _SALDO, usadas), "desc": _columna(titulos, _DESC, usadas),
             "ref": _columna(titulos, _REF, usadas)}
        if c["monto"] is not None or c["deposito"] is not None or c["retiro"] is not None:
            idx, cols = i, c
            break
    if idx is None:
        res.error = ("No encuentro los títulos de la tabla: necesito una columna «Fecha» y el importe "
                     "(«Monto», o «Depósitos» y «Retiros»).")
        return res

    def celda(fila, clave):
        j = cols[clave]
        return fila[j] if j is not None and j < len(fila) else None

    for n, fila in enumerate(filas[idx + 1:], start=idx + 2):
        if not fila or all(v in (None, "") for v in fila):
            continue
        try:
            fecha = lectura.leer_fecha(celda(fila, "fecha"))
            if cols["monto"] is not None:
                monto = lectura.leer_monto(celda(fila, "monto"))
            else:
                dep = lectura.leer_monto(celda(fila, "deposito")) or CERO
                ret = lectura.leer_monto(celda(fila, "retiro")) or CERO
                monto = abs(dep) - abs(ret)
            saldo = lectura.leer_monto(celda(fila, "saldo")) if cols["saldo"] is not None else None
        except ValueError:
            res.ignorados += 1   # renglones de totales, «Saldo anterior», pies de página
            continue
        if fecha is None or not monto:
            res.ignorados += 1
            continue
        res.movimientos.append(Movimiento(
            fila=n, fecha=fecha, descripcion=lectura.leer_texto(celda(fila, "desc"))[:300],
            monto=monto, saldo=saldo, referencia=lectura.leer_texto(celda(fila, "ref"))[:100],
        ))
    if not res.movimientos:
        res.error = "El archivo no trae movimientos que pueda leer."
    return res
